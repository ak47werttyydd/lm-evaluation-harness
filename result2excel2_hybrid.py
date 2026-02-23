import json, glob, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EVAL_DIR = "/home/a84400789/lm-evaluation-harness/eval_results"
EXCEL_DIR = "/home/a84400789/lm-evaluation-harness/results_excel"

# Task -> preferred metric
TASK_METRICS = {
    "arc_challenge": ("acc_norm", "↑"),
    "arc_easy": ("acc_norm", "↑"),
    "hellaswag": ("acc_norm", "↑"),
    "lambada_openai": ("acc", "↑"),
    "openbookqa": ("acc_norm", "↑"),
    "piqa": ("acc_norm", "↑"),
}
TASKS = list(TASK_METRICS.keys())

rows = []
for run_dir in sorted(glob.glob(os.path.join(EVAL_DIR, "*"))):
    if not os.path.isdir(run_dir):
        continue
    run_name = os.path.basename(run_dir)
    # Find the results JSON (nested inside a subfolder)
    jsons = glob.glob(os.path.join(run_dir, "**", "results_*.json"), recursive=True)
    if not jsons:
        continue
    with open(jsons[0]) as f:
        data = json.load(f)

    results = data.get("results", {})
    row = {"name": run_name}
    for task, (metric, _) in TASK_METRICS.items():
        key = f"{metric},none"
        val = results.get(task, {}).get(key, None)
        row[task] = val

    vals = [v for v in [row[t] for t in TASKS] if v is not None]
    row["avg"] = sum(vals) / len(vals) if vals else None

    # Extract gbs number for sorting
    m = re.search(r"gbs(\d+)", run_name)
    row["_sort"] = int(m.group(1)) if m else 0
    rows.append(row)

rows.sort(key=lambda r: r["_sort"])

# Create Excel
wb = Workbook()
ws = wb.active
ws.title = "Eval Results"

header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
header_fill = PatternFill("solid", fgColor="4472C4")
header_align = Alignment(horizontal="center", vertical="center")
data_font = Font(name="Arial", size=11)
avg_font = Font(name="Arial", size=11, bold=True)
avg_fill = PatternFill("solid", fgColor="E2EFDA")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

headers = ["Model"] + TASKS + ["Average"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

for r_idx, row in enumerate(rows, 2):
    c = ws.cell(row=r_idx, column=1, value=row["name"])
    c.font = data_font
    c.border = thin_border

    for c_idx, task in enumerate(TASKS, 2):
        val = row[task]
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.number_format = "0.0000"
        cell.font = data_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    avg_cell = ws.cell(row=r_idx, column=len(TASKS) + 2, value=row["avg"])
    avg_cell.number_format = "0.0000"
    avg_cell.font = avg_font
    avg_cell.fill = avg_fill
    avg_cell.alignment = Alignment(horizontal="center")
    avg_cell.border = thin_border

ws.column_dimensions["A"].width = 50
for col in range(2, len(headers) + 1):
    ws.column_dimensions[chr(64 + col)].width = 18
ws.auto_filter.ref = ws.dimensions

# Auto-increment output filename: eval_results001.xlsx, eval_results002.xlsx, ...
idx = 1
while True:
    out = os.path.join(EXCEL_DIR, f"eval_results{idx:03d}.xlsx")
    if not os.path.exists(out):
        break
    idx += 1

wb.save(out)
print(f"Saved to {out}")